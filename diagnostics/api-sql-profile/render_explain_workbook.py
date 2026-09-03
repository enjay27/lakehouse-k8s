#!/usr/bin/env python3
"""Build the per-case EXPLAIN comparison workbook.

    uv run python render_explain_workbook.py \
        --case admin \
        --matrix reports/doc-api-sql-matrix-admin-20260903-094137.md \
        --explain runs/apiexplain-admin-20260903-110152.json

One workbook per case (Kade, 2026-09-03). Design and the reasoning behind each
sheet are in `PLAN-explain-workbook.md`; this is the implementation.

A VIEW, NOT A SOURCE. Every figure is read from the matrix report and the run
file at build time. `runs/apiexplain-*.json` remains authoritative -- when the
workbook and the run disagree, the run is right and this script has a bug.

NO FORMULAS OVER THE DATA, deliberately. The rows are imported measurements, not
a model: nothing recalculates when a cell changes, because the inputs are JSON
files rather than other cells. The only formulas are the derived counts on
`Provenance`, which really are computed from sheet contents and therefore should
follow them.
"""

import argparse
import collections
import json
import pathlib
import re
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent
REPO = HERE
while not (REPO / "src").is_dir() and REPO != REPO.parent:
    REPO = REPO.parent
sys.path.insert(0, str(REPO / "src"))

from api_report import explain_index  # noqa: E402
from api_trace import REDACTED, extract_verb, normalize_sql  # noqa: E402
from schema_audit import STOCK_INDEX_NAMES  # noqa: E402

FONT = "Arial"
#: Excel's hard limit is 32,767 characters in a cell. A silently truncated cell
#: is worse than a short one, so the marker is visible and the run file is named
#: as the place to go instead.
CELL_MAX = 32_000
HEAD_FILL = PatternFill("solid", fgColor="DDDDDD")
FLAG_FILL = PatternFill("solid", fgColor="FFE0E0")


def truncate(s):
    s = "" if s is None else str(s)
    return s if len(s) <= CELL_MAX else s[:CELL_MAX] + f"… [truncated, {len(s)} chars]"


def where_of(sql):
    m = re.search(r"\bWHERE\b(.*)$", " ".join((sql or "").split()), re.I)
    return m.group(1).strip() if m else ""


def predicate_columns(sql):
    """The SET of columns a statement constrains, order-insensitive.

    THE POINT OF THE `Shapes` SHEET. The predicate column ORDER varies between
    runs, so the same grantee lookup appears under several distinct SQL texts.
    Comparing text reports differences that do not exist; comparing the column
    set does not.
    """
    return tuple(sorted(set(re.findall(r"([a-z_]+)\s*=\s*[?$]", where_of(sql), re.I))))


def plan_cost(e):
    return ((e.get("plan") or {}).get("Plan") or {}).get("Total Cost")


def scan_summary(e):
    """`node on relation using index ×n` -- the same collapse the report uses."""
    scans = e.get("scans") or []
    heap = next((s for s in scans if s.get("relation")), None)
    head = (
        f"{heap.get('node')} on {heap.get('relation')}"
        if heap
        else ", ".join(e.get("node_types") or [])
    )
    counts = {}
    for s in scans:
        if s.get("index"):
            counts[s["index"]] = counts.get(s["index"], 0) + 1
    if counts:
        head += " using " + ", ".join(
            n + (f" ×{c}" if c > 1 else "") for n, c in counts.items()
        )
    return head


def plan_node(e):
    """The top plan node, or {}."""
    return ((e.get("plan") or {}).get("Plan")) or {}


def cond_and_filter(e):
    """`Index Cond` and `Filter` from anywhere in the plan tree.

    THE PAIR IS THE FINDING, at row level. `Index Cond` is what the planner
    pushed INTO the index -- a seek. `Filter` is what it had to test on every
    row it read. The grantee lookup has NO Index Cond and carries its whole
    predicate in Filter, which is what "no usable index" means concretely; the
    securable lookup is the reverse. Sorting on these two columns separates the
    served statements from the unserved without opening a single plan.

    Read from the whole tree: on a Bitmap plan the condition sits on the Bitmap
    Index Scan child, not on the heap node above it.
    """
    conds, filters = [], []

    def walk(n):
        if not isinstance(n, dict):
            return
        if n.get("Index Cond"):
            conds.append(n["Index Cond"])
        if n.get("Filter"):
            filters.append(n["Filter"])
        for c in n.get("Plans") or []:
            walk(c)

    walk(plan_node(e))
    return " AND ".join(conds), " AND ".join(filters)


def parse_matrix(text):
    """Every statement occurrence in a rendered matrix report, in order.

    Returns list of dicts: api, http_status, seq, table, verb, duration_ms,
    sql, params.
    """
    out = []
    for block in re.split(r"^### `", text, flags=re.M)[1:]:
        api = block.split("`")[0]
        st = re.search(r"→ \*\*(\d+)\*\*", block)
        status = int(st.group(1)) if st else None
        for m in re.finditer(
            r"\*\*\[(?P<seq>\d+)\]\*\* `(?P<table>[^`]*)` · (?P<verb>[^·]*) · "
            r"(?P<dur>[^\n]*)\n\n```sql\n(?P<sql>.*?)\n```\n"
            r"(?:params: `(?P<params>.*?)`\n)?",
            block,
            re.S,
        ):
            dur = re.match(r"([\d.]+)\s*ms", (m.group("dur") or "").strip())
            out.append(
                {
                    "api": api,
                    "http_status": status,
                    "seq": int(m.group("seq")),
                    "table": m.group("table").strip("`"),
                    "verb": (m.group("verb") or "").strip(),
                    "duration_ms": float(dur.group(1)) if dur else None,
                    "sql": m.group("sql"),
                    "params": m.group("params"),
                }
            )
    return out


def style(ws, widths, flag_col=None, n_rows=0):
    """Header, freeze, autofilter, widths -- and highlight ONE thing."""
    for c, w in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    if n_rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{n_rows + 1}"
    #: Exactly one conditional rule per sheet. Highlighting everything
    #: highlights nothing.
    if flag_col and n_rows:
        col = get_column_letter(flag_col)
        ws.conditional_formatting.add(
            f"{col}2:{col}{n_rows + 1}",
            CellIsRule(operator="equal", formula=["TRUE"], fill=FLAG_FILL),
        )


def write_rows(ws, header, rows, widths, flag_col=None):
    ws.append(header)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
    style(ws, widths, flag_col=flag_col, n_rows=len(rows))


def build(case, matrix_path, explain_path, dest):
    text = matrix_path.read_text(encoding="utf-8")
    run = json.loads(explain_path.read_text(encoding="utf-8"))

    #: The two refusals notebook 05 makes, for the same reason: a workbook built
    #: from a mismatched pair looks entirely correct.
    if run.get("case") != case:
        sys.exit(f"{explain_path.name} is a '{run.get('case')}' run, not '{case}'")
    if not str(run.get("matrix", "")).endswith(matrix_path.name):
        sys.exit(
            f"{explain_path.name} was taken against {run.get('matrix')}, "
            f"not {matrix_path.name}"
        )

    state = run.get("index_state")
    idx = explain_index(run.get("explains") or [])
    occurrences = parse_matrix(text)
    per_api = run.get("per_api", {})
    vol = run.get("volume") or {}

    wb = Workbook()
    wb.remove(wb.active)

    # ---- Summary -----------------------------------------------------------
    ws = wb.create_sheet("Summary")
    status_by_api = {o["api"]: o["http_status"] for o in occurrences}
    rows = []
    for api, r in sorted(per_api.items()):
        seq = r.get("seq_scanned") or []
        rows.append(
            [
                case,
                api,
                status_by_api.get(api),
                r.get("statements"),
                r.get("planned"),
                r.get("skipped"),
                ", ".join(seq) or "",
                ", ".join(r.get("indexes_used") or []) or "",
                bool(r.get("uses_index_only")),
                "grant_records" in seq,
            ]
        )
    write_rows(
        ws,
        [
            "case",
            "api",
            "http_status",
            "statements",
            "planned",
            "skipped",
            "seq_scanned",
            "indexes_used",
            "uses_index_only",
            "scans_grant_records",
        ],
        rows,
        [13, 42, 12, 11, 9, 9, 22, 46, 15, 20],
        flag_col=10,
    )

    # ---- Statements --------------------------------------------------------
    ws = wb.create_sheet("Statements")
    rows = []
    for o in occurrences:
        e = idx.get((normalize_sql(o["sql"]), o["params"]))
        rows.append(
            [
                case,
                o["api"],
                o["http_status"],
                o["seq"],
                o["table"],
                o["verb"],
                o["duration_ms"],
                scan_summary(e) if e else "",
                ", ".join((e or {}).get("indexes_used") or []) or "",
                bool(e and e.get("seq_scanned")),
                (e or {}).get("plan_rows"),
                plan_node(e).get("Startup Cost") if e else None,
                plan_cost(e) if e else None,
                bool(e and e.get("parallel")),
                truncate(cond_and_filter(e)[0]) if e else "",
                truncate(cond_and_filter(e)[1]) if e else "",
                truncate(" ".join((o["sql"] or "").split()))[:120],
                truncate(o["params"]),
                truncate(" ".join((o["sql"] or "").split())),
            ]
        )
    write_rows(
        ws,
        [
            "case",
            "api",
            "http_status",
            "seq",
            "table",
            "verb",
            "duration_ms",
            "plan",
            "index_used",
            "seq_scan",
            "plan_rows",
            "startup_cost",
            "total_cost",
            "parallel",
            "index_cond",
            "filter",
            "sql_short",
            "params",
            "sql_full",
        ],
        rows,
        [13, 40, 11, 6, 22, 8, 12, 46, 30, 10, 11, 12, 11, 10, 56, 56, 60, 34, 60],
        flag_col=10,
    )

    # ---- Distinct ----------------------------------------------------------
    ws = wb.create_sheet("Distinct")
    rows = []
    for e in run.get("explains") or []:
        refused = e.get("error") or e.get("skipped")
        rows.append(
            [
                case,
                e.get("table"),
                e.get("verb"),
                e.get("occurrences"),
                len(e.get("apis") or {}),
                scan_summary(e) if not refused else "",
                ", ".join(e.get("indexes_used") or []) or "",
                bool(e.get("seq_scanned")),
                e.get("plan_rows"),
                plan_node(e).get("Startup Cost"),
                plan_cost(e),
                bool(e.get("parallel")),
                truncate(cond_and_filter(e)[0]),
                truncate(cond_and_filter(e)[1]),
                truncate(", ".join(sorted(e.get("apis") or {}))),
                bool(e.get("params_observed")),
                truncate(where_of(e.get("sql"))),
                truncate(e.get("params")),
                not bool(refused),
                truncate(refused) if refused else "",
                truncate(" ".join((e.get("sql") or "").split())),
            ]
        )
    write_rows(
        ws,
        [
            "case",
            "table",
            "verb",
            "occurrences",
            "api_count",
            "plan",
            "index_used",
            "seq_scan",
            "plan_rows",
            "startup_cost",
            "total_cost",
            "parallel",
            "index_cond",
            "filter",
            "api_list",
            "params_observed",
            "where_clause",
            "params",
            "replayable",
            "refusal_reason",
            "sql_full",
        ],
        rows,
        [
            13,
            24,
            8,
            12,
            10,
            46,
            30,
            10,
            11,
            12,
            11,
            10,
            56,
            56,
            70,
            16,
            60,
            34,
            11,
            40,
            60,
        ],
        flag_col=8,
    )

    # ---- Shapes ------------------------------------------------------------
    ws = wb.create_sheet("Shapes")
    shapes = {}
    for e in run.get("explains") or []:
        key = (e.get("verb"), e.get("table"), predicate_columns(e.get("sql")))
        s = shapes.setdefault(
            key,
            {
                "texts": set(),
                "occ": 0,
                "apis": set(),
                "plan": "",
                "idx": "",
                "seq": False,
            },
        )
        s["texts"].add(normalize_sql(e.get("sql") or ""))
        s["occ"] += e.get("occurrences") or 0
        s["apis"].update(e.get("apis") or {})
        if not (e.get("error") or e.get("skipped")):
            s["plan"] = s["plan"] or scan_summary(e)
            s["idx"] = s["idx"] or ", ".join(e.get("indexes_used") or [])
            s["seq"] = s["seq"] or bool(e.get("seq_scanned"))
    rows = [
        [
            case,
            v,
            t,
            ", ".join(cols) or "(none)",
            len(s["texts"]),
            s["occ"],
            len(s["apis"]),
            s["plan"],
            s["idx"],
            s["seq"],
        ]
        for (v, t, cols), s in sorted(
            shapes.items(), key=lambda kv: (-kv[1]["occ"], str(kv[0]))
        )
    ]
    write_rows(
        ws,
        [
            "case",
            "verb",
            "table",
            "predicate_columns",
            "texts",
            "occurrences",
            "api_count",
            "plan",
            "index_used",
            "seq_scan",
        ],
        rows,
        [13, 8, 24, 62, 8, 13, 10, 46, 30, 10],
        flag_col=10,
    )

    # ---- Unplanned ---------------------------------------------------------
    ws = wb.create_sheet("Unplanned")
    rows = []
    for e in run.get("explains") or []:
        why = e.get("error") or e.get("skipped")
        if not why:
            continue
        kind = "redacted" if REDACTED in str(e.get("params")) else "refused"
        rows.append(
            [
                case,
                e.get("table"),
                e.get("verb"),
                kind,
                truncate(str(why)),
                bool(e.get("params_observed")),
                truncate(e.get("params")),
                truncate(" ".join((e.get("sql") or "").split())),
            ]
        )
    #: Statements in the report that never reached the worklist at all.
    for o in occurrences:
        if (normalize_sql(o["sql"]), o["params"]) in idx:
            continue
        kind = "redacted" if o["params"] and REDACTED in o["params"] else "unmatched"
        rows.append(
            [
                case,
                o["table"],
                o["verb"],
                kind,
                "not in the sweep's worklist",
                bool(o["params"]),
                truncate(o["params"]),
                truncate(" ".join((o["sql"] or "").split())),
            ]
        )
    write_rows(
        ws,
        [
            "case",
            "table",
            "verb",
            "kind",
            "reason",
            "params_observed",
            "params",
            "sql_full",
        ],
        rows,
        [13, 30, 8, 12, 46, 16, 34, 70],
    )

    # ---- Schema ------------------------------------------------------------
    ws = wb.create_sheet("Schema")
    rows = [
        [
            i.get("name"),
            run.get("indexes_on_focus_table") and "grant_records",
            truncate(" ".join((i.get("def") or "").split())),
            i.get("name") in STOCK_INDEX_NAMES,
        ]
        for i in (run.get("indexes_on_focus_table") or [])
    ]
    write_rows(
        ws,
        ["index_name", "table", "definition", "in_upstream_schema"],
        rows,
        [34, 18, 96, 20],
    )

    # ---- Provenance --------------------------------------------------------
    ws = wb.create_sheet("Provenance")
    n = len(per_api)
    prov = [
        ["case", case],
        ["matrix_report", matrix_path.name],
        ["explain_run", explain_path.name],
        ["index_state", state],
        ["idx_grant_records_grantee", "absent" if state == "absent" else "present"],
        ["", ""],
        ["entities rows", vol.get("entities")],
        ["grant_records rows", vol.get("grant_records")],
        ["policy_mapping_record rows", vol.get("policy_mapping_record")],
        [
            "principal_authentication_data rows",
            vol.get("principal_authentication_data"),
        ],
        ["", ""],
        ["APIs", n],
        ["statement occurrences", len(occurrences)],
        ["distinct (SQL, params) pairs", run.get("pairs")],
        ["APIs that seq-scan grant_records", f"=COUNTIF(Summary!J2:J{n + 1},TRUE)"],
        ["", ""],
        ["READ THIS", ""],
        [
            "total_cost / plan_rows are planner ESTIMATES",
            "They move with ANALYZE. Two sweeps 25 min apart differed in 11 of 264 "
            "plan_rows while 0 of 264 differed in plan SHAPE. Do not average, sum "
            "or trend them — sort and compare shape.",
        ],
        [
            "duration_ms is NOT from the EXPLAIN",
            "Plain EXPLAIN does not execute. It is the drive's observed "
            "server-side duration, a different measurement; the same plan has "
            "measured a 4.6x spread.",
        ],
        [
            "one volume",
            "Every plan is relative to the row counts above. A larger realm is a "
            "different regime, not more of this one.",
        ],
        [
            "authority",
            f"{explain_path.name} and {matrix_path.name}. Where this workbook and "
            "the run disagree, the run is right.",
        ],
    ]
    for r in prov:
        ws.append(r)
    for row in ws.iter_rows():
        row[0].font = Font(name=FONT, bold=True)
        for c in row[1:]:
            c.font = Font(name=FONT)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"] = ws.column_dimensions["A"]
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 96

    dest.parent.mkdir(exist_ok=True)
    wb.save(dest)
    return {
        "summary": n,
        "statements": len(occurrences),
        "distinct": len(run.get("explains") or []),
        "shapes": len(shapes),
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--case", required=True)
    ap.add_argument("--matrix", required=True)
    ap.add_argument("--explain", required=True)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    from datetime import datetime

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = pathlib.Path(
        args.out or HERE / "reports" / f"api-explain-{args.case}-{stamp}.xlsx"
    )
    counts = build(args.case, HERE / args.matrix, HERE / args.explain, dest)
    print(f"-> {dest.name}")
    for k, v in counts.items():
        print(f"   {k}: {v}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
