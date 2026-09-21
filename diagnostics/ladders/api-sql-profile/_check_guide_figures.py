#!/usr/bin/env python3
"""Check every figure the Korean guide/results docs quote against its source.

    python3 _check_guide_figures.py

WHY THIS EXISTS. `render_index_findings.py` and `render_explain_workbook.py` are
generators -- their numbers cannot drift because nobody types them. The two
Korean documents are PROSE, written by hand, and prose drifts from the runs that
produced it without anyone noticing until someone tries to reproduce it. This
script is the substitute for generating them: it recomputes each quoted figure
from the workbooks and the run files, then asserts the document still says it.

DIRECTION OF THE CHECK, WHICH IS THE POINT. A figure whose anchor pattern does
not match at all is reported as NOT QUOTED and does not fail -- the document is
allowed to stop making a claim. A figure whose anchor DOES match with a
different value FAILS. So editing prose is cheap; silently editing a number is
not.

The workbook is a view. Where it and `runs/apiexplain-*.json` disagree the run
file is right, so §0 asserts they agree before anything else is checked.
"""

import json
import pathlib
import re
import sys

from openpyxl import load_workbook

HERE = pathlib.Path(__file__).resolve().parent
REPORTS = HERE / "reports"
RUNS = HERE / "runs"
GUIDE = HERE / "doc-api-sql-profile-guide-ko.md"
RESULTS = HERE / "doc-api-sql-profile-results-ko.md"
CASES = ("unauthorized", "authorized", "admin")

#: The eight sheets the workbook ships today. `PLAN-explain-workbook.md` and
#: `README.md` both still say seven -- `Cost` was added afterwards -- and the
#: guide says so out loud, so the count is checked rather than trusted.
SHEETS = (
    "Summary",
    "Statements",
    "Distinct",
    "Cost",
    "Shapes",
    "Unplanned",
    "Schema",
    "Provenance",
)


def sheet(wb, name):
    it = wb[name].iter_rows(values_only=True)
    head = list(next(it))
    return [dict(zip(head, r)) for r in it]


def provenance(wb):
    return {r[0]: r[1] for r in wb["Provenance"].iter_rows(values_only=True) if r[0]}


def load(case):
    wb = load_workbook(REPORTS / f"api-explain-{case}-latest.xlsx", data_only=True)
    return wb, {n: sheet(wb, n) for n in SHEETS if n != "Provenance"}, provenance(wb)


def is_grantee(row):
    """The grantee lookup, identified by its predicate COLUMN SET.

    Not by SQL text: `predicate_columns` on the `Shapes` sheet exists because
    the column ORDER varies between runs, so the same lookup appears under
    several distinct texts. Matching text here would miss rows at random.
    """
    cols = set(re.findall(r"[a-z_]+", (row.get("where_clause") or "").lower()))
    return "grantee_id" in cols and "securable_id" not in cols


def main():
    figures = {}  # label -> (expected, anchor regex, which doc)
    problems = []

    books = {}
    for case in CASES:
        wb, sh, prov = load(case)
        books[case] = (sh, prov)

        # --- 0. the workbook is a view; the run file is the authority --------
        run_name = prov["explain_run"]
        run_path = RUNS / run_name
        if not run_path.exists():
            problems.append(
                f"{case}: Provenance names {run_name}, which is not in runs/"
            )
            continue
        run = json.loads(run_path.read_text())
        if run["case"] != case:
            problems.append(f"{case}: run file says case={run['case']!r}")
        if pathlib.Path(run["matrix"]).name != prov["matrix_report"]:
            problems.append(
                f"{case}: Provenance matrix_report {prov['matrix_report']!r} != "
                f"run matrix {run['matrix']!r}"
            )
        for key, col in (
            ("apis", "APIs"),
            ("instances", "statement occurrences"),
            ("pairs", "distinct (SQL, params) pairs"),
        ):
            if int(run[key]) != int(prov[col]):
                problems.append(
                    f"{case}: run {key}={run[key]} != Provenance {col}={prov[col]}"
                )
        for table, rows in run["volume"].items():
            if int(prov[f"{table} rows"]) != int(rows):
                problems.append(f"{case}: volume {table} disagrees with Provenance")

        missing = [n for n in SHEETS if n not in wb.sheetnames]
        if missing:
            problems.append(f"{case}: workbook is missing sheets {missing}")

    if problems:
        print("PROVENANCE FAILED -- nothing else was checked:")
        for p in problems:
            print("  ", p)
        return 1

    # --- 1. the headline ----------------------------------------------------
    per_case_scanning = {
        case: sum(
            1
            for r in books[case][0]["Summary"]
            if str(r["scans_grant_records"]) == "True"
        )
        for case in CASES
    }
    if len(set(per_case_scanning.values())) != 1:
        problems.append(
            f"cases no longer agree on scans_grant_records: {per_case_scanning}"
        )
    total_pairs = sum(len(books[case][0]["Summary"]) for case in CASES)
    figures["headline: APIs scanning grant_records, per case"] = (
        f"{per_case_scanning['admin']}",
        r"scans_grant_records`가 세 파일 모두에서 (\d+)/43 TRUE",
        RESULTS,
    )
    figures["headline: total (case, API) pairs"] = (
        f"{total_pairs}",
        r"신원 = (\d+)개 \(케이스, API\) 쌍 전부가",
        RESULTS,
    )
    fis = sum(
        1
        for case in CASES
        for r in books[case][0]["Summary"]
        if str(r["fully_index_served"]) == "True"
    )
    if fis:
        problems.append(
            f"fully_index_served is TRUE for {fis} rows; the docs say none are"
        )

    # --- 2. per-case scale table -------------------------------------------
    for case in CASES:
        sh, prov = books[case]
        row = (
            rf"\|\s*`{case}`\s*\|\s*43\s*\|\s*(\d+)\s*\|\s*(\d+)\s*\|\s*(\d+)"
            rf"\s*\|\s*(\d+)\s*\|\s*(\d+)\s*\|"
        )
        expected = (
            len(sh["Statements"]),
            len(sh["Distinct"]),
            len(sh["Cost"]),
            len(sh["Shapes"]),
            len(sh["Unplanned"]),
        )
        figures[f"scale row: {case}"] = (
            " ".join(str(v) for v in expected),
            row,
            RESULTS,
        )

    # --- 3. volume ----------------------------------------------------------
    prov = books["admin"][1]
    for table, label in (
        ("entities", "entities"),
        ("grant_records", "grant_records"),
        ("policy_mapping_record", "policy_mapping_record"),
        ("principal_authentication_data", "principal_authentication_data"),
    ):
        n = int(prov[f"{table} rows"])
        figures[f"volume: {label}"] = (
            f"{n:,}",
            rf"\|\s*`{re.escape(label)}`\s*\|\s*([\d,]+)\s*\|",
            RESULTS,
        )

    # --- 4. the two lookups -------------------------------------------------
    dist = books["admin"][0]["Distinct"]
    grantee = [
        r
        for r in dist
        if r["table"] == "grant_records" and r["verb"] == "SELECT" and is_grantee(r)
    ]
    securable = [
        r
        for r in dist
        if r["table"] == "grant_records"
        and r["verb"] == "SELECT"
        and str(r["plan"]).startswith("Index Only Scan")
    ]
    deletes = [
        r for r in dist if r["table"] == "grant_records" and r["verb"] == "DELETE"
    ]
    inserts = [
        r for r in dist if r["table"] == "grant_records" and r["verb"] == "INSERT"
    ]
    if not (grantee and securable and deletes and inserts):
        problems.append("could not locate all four grant_records shapes on Distinct")
    else:
        figures["grantee SELECT total_cost"] = (
            f"{grantee[0]['total_cost']}",
            r"\*\*Seq Scan\*\*\s*\|\s*—\s*\|\s*1\s*\|\s*([\d.]+)\s*\|",
            RESULTS,
        )
        figures["securable SELECT total_cost"] = (
            f"{securable[0]['total_cost']}",
            r"Index Only Scan\s*\|\s*`grant_records_pkey`\s*\|\s*20\s*\|\s*([\d.]+)\s*\|",
            RESULTS,
        )
        figures["DELETE total_cost"] = (
            f"{deletes[0]['total_cost']}",
            r"ModifyTable, \*\*Seq Scan\*\*\s*\|\s*—\s*\|\s*1\s*\|\s*([\d.]+)\s*\|",
            RESULTS,
        )
        figures["INSERT total_cost"] = (
            f"{inserts[0]['total_cost']}",
            r"ModifyTable, Result\s*\|\s*—\s*\|\s*0\s*\|\s*([\d.]+)\s*\|",
            RESULTS,
        )
        figures["grantee rows_scanned"] = (
            f"{int(grantee[0]['rows_scanned']):,}",
            r"플래너는 한 행이 나올 것을 예상하면서 ([\d,]+)행을 읽습니다",
            RESULTS,
        )

    # --- 5. cost concentration ---------------------------------------------
    for case in CASES:
        cost = books[case][0]["Cost"]
        total = sum(float(r["cost_units"] or 0) for r in cost) or 1.0
        seq = sum(
            float(r["cost_units"] or 0)
            for r in cost
            if r["table"] == "grant_records" and str(r["plan"]).startswith("Seq Scan")
        )
        top = max(
            (r for r in cost if r["table"] == "grant_records" and is_grantee(r)),
            key=lambda r: float(r["cost_units"] or 0),
        )
        figures[f"cost share: {case}"] = (
            f"{seq / total * 100:.1f} {float(top['cost_share']):.4f} {int(top['occurrences'])}",
            rf"\|\s*`{case}`\s*\|\s*\*\*([\d.]+) %\*\*\s*\|\s*([\d.]+) \((\d+)회",
            RESULTS,
        )

    # --- 6. outcomes, and that refusals reach the scan ----------------------
    for case in CASES:
        summ = books[case][0]["Summary"]
        st = [str(r["http_status"]) for r in summ]
        ok = sum(1 for s in st if s.startswith("2"))
        f403 = st.count("403")
        f404 = st.count("404")
        dist_c = books[case][0]["Distinct"]
        lookups = max(
            int(r["occurrences"])
            for r in dist_c
            if r["table"] == "grant_records" and is_grantee(r)
        )
        figures[f"outcomes: {case}"] = (
            f"{ok} {f403} {f404} {lookups}",
            rf"\|\s*`{case}`\s*\|\s*(\d+)\s*\|\s*(\d+)\s*\|\s*(\d+)\s*\|\s*(\d+)\s*\|\s*43 / 43\s*\|",
            RESULTS,
        )

    # --- 7. the contrast: entities is never sequentially scanned ------------
    ent_seq = sum(
        1
        for case in CASES
        for r in books[case][0]["Distinct"]
        if r["table"] == "entities" and str(r["seq_scan"]) == "True"
    )
    figures["entities rows scanned sequentially"] = (
        f"{ent_seq}",
        r"\*\*`entities`를 순차 스캔하는 행은 (\d+)개\*\*",
        RESULTS,
    )

    # --- 8. the schema inventory -------------------------------------------
    schema = books["admin"][0]["Schema"]
    if len(schema) != 1 or schema[0]["index_name"] != "grant_records_pkey":
        problems.append(
            f"Schema sheet no longer holds only grant_records_pkey: {schema}"
        )
    if str(schema[0]["in_upstream_schema"]) != "True":
        problems.append("grant_records_pkey is no longer marked in_upstream_schema")
    if books["admin"][1]["index_state"] != "absent":
        problems.append("index_state is no longer 'absent'; the results doc says it is")

    # --- 9. the 194x claim comes from the workbook itself ------------------
    read_this = " ".join(str(v) for v in books["admin"][1].values() if v)
    if "194x" not in read_this:
        problems.append(
            "Provenance no longer carries the 194x note the results doc quotes"
        )

    # --- 10. sheet count the guide states ----------------------------------
    #: The guide spells the count in Korean, so the expectation is spelled too --
    #: comparing "8" to "여덟" would fail on a document that is perfectly correct.
    figures["sheet count"] = (
        {7: "일곱", 8: "여덟"}[len(SHEETS)],
        r"현재는 \*\*(여덟|일곱) 개\*\*입니다",
        GUIDE,
    )

    # ---- report ------------------------------------------------------------
    checked = quoted = 0
    for label, (expected, pattern, doc) in figures.items():
        text = doc.read_text(encoding="utf-8")
        m = re.search(pattern, text)
        checked += 1
        if not m:
            print(f"  NOT QUOTED  {label}  (expected {expected})")
            continue
        quoted += 1
        got = " ".join(g for g in m.groups() if g is not None)
        got_norm = got.replace(" %", "")
        if got_norm != expected and got_norm.replace(",", "") != expected.replace(
            ",", ""
        ):
            problems.append(
                f"{doc.name}: {label} says {got!r}, source says {expected!r}"
            )

    print(f"\n{quoted}/{checked} figures quoted and compared.")
    if problems:
        print("\nFAILED:")
        for p in problems:
            print("  ", p)
        return 1
    print("All quoted figures match their sources.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
